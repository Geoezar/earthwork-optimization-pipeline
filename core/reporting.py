from __future__ import annotations

import math
from pathlib import Path

import pandas as pd
import qrcode
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from core.optimization import OptimizationResult


DEFAULT_JURY_RESULTS_URL = "https://github.com/Geoezar/earthwork-optimization-pipeline"


def _set_apa_table_style(table) -> None:
    table.style = "Light Shading"
    table.autofit = True


def _add_apa_paragraph(document: Document, text: str, indent: bool = True, bold: bool = False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 2.0
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    run = paragraph.add_run(text)
    run.bold = bold
    return paragraph


def _add_apa_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = True


def _add_table_caption(document: Document, number: int, title: str) -> None:
    document.add_paragraph(f"Table {number}").runs[0].bold = True
    document.add_paragraph(title).runs[0].italic = True


def create_word_report(
    result: OptimizationResult,
    frame: pd.DataFrame,
    output_directory: Path,
    all_results: dict[str, OptimizationResult] | None = None,
    summary_png_path: Path | None = None,
    jury_results_url: str = DEFAULT_JURY_RESULTS_URL,
) -> Path:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    for _ in range(4):
        document.add_paragraph()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 2.0
    title_run = title.add_run(f"Topographic Optimization and Earthwork Cost Analysis Report: {result.area_name}")
    title_run.bold = True

    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.line_spacing = 2.0
    author.add_run("Earthwork Optimization Pipeline\n")
    author.add_run("Geological Engineering Design Pipeline\n")
    author.add_run("Machine-Learning-Assisted Earthwork Optimization\n")
    document.add_page_break()

    _add_apa_heading(document, "Optimization Decision Mechanism", level=1)
    if all_results:
        _add_apa_paragraph(
            document,
            "All candidate areas were evaluated under the maximum 2.0% design slope constraint. "
            "The decision criterion is the minimum total earthwork volume after optimization.",
        )
        for area_name, area_result in all_results.items():
            _add_apa_paragraph(
                document,
                f"{area_name}: {area_result.total_earthwork_m3:,.2f} m3 total earthwork.",
            )
        _add_apa_paragraph(
            document,
            f"{result.area_name} is selected as the optimum area because it has the lowest total earthwork volume.",
            bold=True,
        )
    else:
        _add_apa_paragraph(document, "This report was generated as a single-area engineering analysis.")

    _add_apa_heading(document, "Volume and Quantity Results", level=1)
    _add_apa_paragraph(
        document,
        f"The natural slope of {result.area_name} is {result.original_slope_percent:.2f}%. "
        f"The design slope after centroid-locked optimization is {result.design_slope_percent:.2f}%. "
        "Volumes were calculated by grid-cell integration against the selected RBF terrain surface.",
    )

    _add_table_caption(document, 1, f"{result.area_name} Optimization Volume Results")
    volume_table = document.add_table(rows=1, cols=2)
    _set_apa_table_style(volume_table)
    volume_table.rows[0].cells[0].text = "Parameter"
    volume_table.rows[0].cells[1].text = "Value"
    volume_rows = [
        ("Original natural slope", f"{result.original_slope_percent:.4f}%"),
        ("Design slope", f"{result.design_slope_percent:.4f}%"),
        ("Cut volume", f"{result.cut_volume_m3:,.2f} m3"),
        ("Fill volume", f"{result.fill_volume_m3:,.2f} m3"),
        ("Total earthwork", f"{result.total_earthwork_m3:,.2f} m3"),
        ("RBF model", f"{result.rbf_model} (RMSE: {result.rbf_rmse:.3f} m)"),
    ]
    for label, value in volume_rows:
        row = volume_table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value

    balance = document.add_paragraph()
    balance.paragraph_format.line_spacing = 2.0
    balance.add_run("Material balance: ").bold = True
    balance.add_run(result.balance_warning)

    _add_apa_heading(document, "Geological Orientation and Stability", level=1)
    _add_apa_paragraph(
        document,
        "The natural trend plane and optimized design plane were converted into structural geology notation "
        "for field interpretation.",
    )

    _add_table_caption(document, 2, "Geological Orientation and Stability Comparison")
    geology_table = document.add_table(rows=1, cols=3)
    _set_apa_table_style(geology_table)
    geology_table.rows[0].cells[0].text = "Parameter"
    geology_table.rows[0].cells[1].text = "Natural condition"
    geology_table.rows[0].cells[2].text = "Design condition"
    geology_rows = [
        ("Azimuth", result.original_azimuth, result.design_azimuth),
        ("Quadrant", result.original_quadrant, result.design_quadrant),
        ("Slope", f"{result.original_slope_percent:.2f}%", f"{result.design_slope_percent:.2f}%"),
    ]
    for label, natural, design in geology_rows:
        row = geology_table.add_row()
        row.cells[0].text = label
        row.cells[1].text = natural
        row.cells[2].text = design

    _add_apa_paragraph(document, f"Geotechnical interpretation: {result.geotechnical_warning}", bold=True)

    natural_angle = math.degrees(math.atan(result.original_slope_percent / 100))
    design_angle = math.degrees(math.atan(result.design_slope_percent / 100))
    angle_change = abs(natural_angle - design_angle)
    proof = document.add_paragraph(
        f"The natural slope angle is approximately {natural_angle:.2f} degrees, "
        f"while the optimized design slope angle is approximately {design_angle:.2f} degrees. "
        f"The angular change is {angle_change:.2f} degrees, but it controls "
        f"{result.total_earthwork_m3:,.2f} m3 of earth movement."
    )
    proof.paragraph_format.line_spacing = 2.0
    for run in proof.runs:
        run.italic = True

    _add_apa_heading(document, "Trend Surface Matrix", level=1)
    _add_apa_paragraph(
        document,
        "The baseline trend plane was solved as a three-unknown least-squares system. "
        "The matrix terms below provide auditability for the fitted plane.",
    )

    matrix_table = document.add_table(rows=3, cols=7)
    rows = matrix_table.rows
    rows[0].cells[0].text = f"{result.regression.n}"
    rows[0].cells[1].text = f"{result.regression.sum_x1:,.1f}"
    rows[0].cells[2].text = f"{result.regression.sum_x2:,.1f}"
    rows[0].cells[3].text = "x"
    rows[0].cells[4].text = "a"
    rows[0].cells[5].text = "="
    rows[0].cells[6].text = f"{result.regression.sum_y:,.1f}"
    rows[1].cells[0].text = f"{result.regression.sum_x1:,.1f}"
    rows[1].cells[1].text = f"{result.regression.sum_x1_sq:,.1f}"
    rows[1].cells[2].text = f"{result.regression.sum_x1_x2:,.1f}"
    rows[1].cells[3].text = "x"
    rows[1].cells[4].text = "b"
    rows[1].cells[5].text = "="
    rows[1].cells[6].text = f"{result.regression.sum_x1_y:,.1f}"
    rows[2].cells[0].text = f"{result.regression.sum_x2:,.1f}"
    rows[2].cells[1].text = f"{result.regression.sum_x1_x2:,.1f}"
    rows[2].cells[2].text = f"{result.regression.sum_x2_sq:,.1f}"
    rows[2].cells[3].text = "x"
    rows[2].cells[4].text = "c"
    rows[2].cells[5].text = "="
    rows[2].cells[6].text = f"{result.regression.sum_x2_y:,.1f}"

    _add_apa_paragraph(
        document,
        f"Solved coefficients: a = {result.regression.intercept:.5f}, "
        f"b = {result.regression.x1_coefficient:.5f}, "
        f"c = {result.regression.x2_coefficient:.5f}.",
        bold=True,
    )

    _add_apa_heading(document, "Cost Analysis", level=1)
    _add_apa_paragraph(
        document,
        "The budget uses official 2026 unit prices for excavation, fill placement, loading, and KGM transport. "
        "The published unit prices already include contractor overhead and profit.",
    )

    _add_table_caption(document, 3, "Official Earthwork Cost Breakdown")
    cost_table = document.add_table(rows=1, cols=2)
    _set_apa_table_style(cost_table)
    cost_table.rows[0].cells[0].text = "Cost item"
    cost_table.rows[0].cells[1].text = "Amount"
    cost = result.cost_breakdown
    cost_rows = [
        ("Machine excavation", f"TL {cost.excavation_cost:,.2f}"),
        ("Fill spreading and compaction", f"TL {cost.fill_cost:,.2f}"),
        ("Truck loading and unloading", f"TL {cost.loading_cost:,.2f}"),
        ("External transport", f"TL {cost.transport_cost:,.2f}"),
        ("Total project budget", f"TL {cost.total_cost:,.2f}"),
    ]
    for label, value in cost_rows:
        row = cost_table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value

    if result.sensitivity:
        _add_apa_heading(document, "Machine Learning Cost Sensitivity", level=1)
        saving = result.cost_breakdown.total_cost - result.sensitivity.theoretical_minimum_cost
        _add_apa_paragraph(
            document,
            f"The sensitivity scan identified a theoretical minimum at "
            f"{result.sensitivity.theoretical_optimum_slope:.2f}% design slope with an estimated cost of "
            f"TL {result.sensitivity.theoretical_minimum_cost:,.2f}. "
            f"The difference against the selected design budget is TL {saving:,.2f}.",
            bold=abs(saving) > 0,
        )

    if result.mass_haul:
        _add_apa_heading(document, "Internal Mass-Haul Routing", level=1)
        _add_apa_paragraph(
            document,
            f"The greedy nearest-fill routing model matched {result.mass_haul.total_internal_volume_m3:,.2f} m3 "
            f"of material inside the site. The weighted average internal movement distance is "
            f"{result.mass_haul.average_distance_m:.2f} m.",
        )

    if summary_png_path and summary_png_path.exists():
        document.add_page_break()
        _add_apa_heading(document, "Visual Summary", level=1)
        document.add_paragraph("Figure 1").runs[0].bold = True
        document.add_paragraph(f"{result.area_name} terrain model, volume distribution, and cut/fill map").runs[0].italic = True
        document.add_picture(str(summary_png_path), width=Inches(6.2))
        document.add_paragraph("Red regions indicate cut zones; blue regions indicate fill zones.")

    _add_apa_heading(document, "Digital Jury Access", level=1)
    _add_apa_paragraph(
        document,
        "Scan the QR code to open the GitHub-rendered jury results page. The page summarizes the full project, "
        "lists the generated outputs, and explains how to download interactive HTML artifacts from the repository.",
    )

    document.add_paragraph("Figure 2").runs[0].bold = True
    document.add_paragraph("GitHub Jury Results QR Code").runs[0].italic = True
    qr_path = output_directory / f"{result.area_name.replace(' ', '_')}_QR.png"
    qrcode.make(jury_results_url).save(str(qr_path))
    document.add_picture(str(qr_path), width=Inches(1.5))

    output_path = output_directory / f"{result.area_name.replace(' ', '_')}_APA7_Report.docx"
    document.save(str(output_path))
    return output_path

